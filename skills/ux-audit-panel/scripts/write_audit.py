#!/usr/bin/env python3
"""Fill the audit template from merged evaluator results.

usage: write_audit.py <results.json> <output.xlsx> [template.xlsx]
"""
import json
import math
import shutil
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

EVAL_COLS = ["C", "D", "J", "K", "L"]
NOTES_WIDTH = 72
SUMMARY_ROW = {
    "1. Consistency": 3, "2. Aesthetic": 4, "3.Recover": 5, "4. Recognition": 6,
    "5. Accessibility": 7, "6. Attention": 8, "7. Status": 9, "8. Real world": 10,
    "9. Important": 11,
}
SHEETS = {
    "1. Consistency": (4, 9),
    "2. Aesthetic": (4, 10),
    "3.Recover": (4, 8),
    "4. Recognition": (4, 7),
    "5. Accessibility": (4, 9),
    "6. Attention": (4, 9),
    "7. Status": (4, 8),
    "8. Real world": (4, 7),
    "9. Important": (4, 9),
}


INSTRUCTIONS = [
    "Hi \U0001F44B",
    "This audit was produced by the ux-audit-panel skill. Several evaluators scored the "
    "same 49 statements independently, each from a different design perspective, without "
    "seeing each other's work or scores.",
    "How to read it:",
    "1. Every evaluator has their own score column, headed by the perspective they audited "
    "from. Their scores are averaged in the Average column, which drives the percentages "
    "and the Summary sheet.",
    "2. The Notes column holds each evaluator's reasoning in their own words, labelled with "
    "the perspective that wrote it. Where they read the same screen differently, you will "
    "see it there.",
    "3. A statement marked N/A had no evidence to judge it against. It is left out of the "
    "score rather than counted as a zero.",
    "4. Don't move the cells that say \"don't move\".",
    "5. The statements where the evaluators disagreed most are written up in "
    "divergence-report.md, saved next to this file. Read that one first.",
]


def fill_instructions(wb):
    ws = wb["Instructions"]
    template_style = ws["C4"]._style
    for offset, text in enumerate(INSTRUCTIONS):
        cell = ws.cell(row=3 + offset, column=3)
        cell.value = text
        cell._style = copy(template_style)
        ws.row_dimensions[3 + offset].height = estimate_height(text, 35)


def estimate_height(text, width_chars, line_px=13.5, minimum=18.75):
    """Rough row height for wrapped text in a column this many characters wide."""
    lines = 0
    for paragraph in str(text).split("\n"):
        lines += max(1, math.ceil(len(paragraph) / width_chars))
    return min(409, max(minimum, lines * line_px + 8))


def style_evaluator_column(ws, letter, source_letter, first, last):
    """Give a new evaluator column the same look as the two the template ships with."""
    ws.column_dimensions[letter].width = ws.column_dimensions[source_letter].width
    for row in [3] + list(range(first, last + 1)):
        target, source = ws[f"{letter}{row}"], ws[f"{source_letter}{row}"]
        target._style = copy(source._style)


def fill_project(wb, project):
    ws = wb["About the project"]
    for cell, key in (("D3", "app_name"), ("D4", "flow"), ("D5", "goal"),
                      ("D6", "link"), ("D7", "notes")):
        if project.get(key):
            ws[cell] = project[key]


def wrap_iferror(cell):
    """Make an AVERAGE formula survive a range with nothing numeric in it."""
    formula = cell.value
    if formula and formula.startswith("=") and "IFERROR" not in formula:
        cell.value = f'=IFERROR({formula[1:]},"N/A")'


def fill_sheet(ws, first, last, rows, evaluators):
    cols = EVAL_COLS[:len(evaluators)]
    for col in cols[2:]:
        style_evaluator_column(ws, col, "D", first, last)
    for col, name in zip(cols, evaluators):
        ws[f"{col}3"] = name
    ws.column_dimensions["I"].width = NOTES_WIDTH
    for row in range(first, last + 1):
        entry = rows.get(str(row)) or rows.get(row)
        if entry is None:
            raise SystemExit(f"{ws.title}: no result for row {row}")
        scores = entry["scores"]
        if len(scores) != len(evaluators):
            raise SystemExit(f"{ws.title} row {row}: {len(scores)} scores for "
                             f"{len(evaluators)} evaluators")
        numeric = []
        for col, score in zip(cols, scores):
            if score is None or score == "N/A":
                ws[f"{col}{row}"] = "N/A"
            else:
                ws[f"{col}{row}"] = float(score)
                numeric.append(f"{col}{row}")
        note = entry.get("note") or ""
        note_cell = ws[f"I{row}"]
        note_cell.value = note
        note_cell.alignment = Alignment(wrap_text=True, horizontal="left",
                                        vertical="top")
        ws.row_dimensions[row].height = max(
            ws.row_dimensions[row].height or 0,
            estimate_height(note, NOTES_WIDTH - 2),
        )
        if numeric:
            ws[f"E{row}"] = f"=ROUND(AVERAGE({','.join(numeric)}),2)"
        else:
            ws[f"E{row}"] = None
            ws[f"G{row}"] = None
            ws[f"H{row}"] = None
    return any(ws[f"E{r}"].value is not None for r in range(first, last + 1))


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    results = json.loads(Path(sys.argv[1]).read_text())
    output = Path(sys.argv[2])
    template = Path(sys.argv[3]) if len(sys.argv) > 3 else \
        Path(__file__).resolve().parent.parent / "assets" / "template.xlsx"

    evaluators = results["evaluators"]
    if not 2 <= len(evaluators) <= 5:
        raise SystemExit("evaluators must be between 2 and 5")

    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, output)
    wb = openpyxl.load_workbook(output)
    fill_project(wb, results.get("project", {}))
    fill_instructions(wb)
    summary = wb["Summary (global results)"]
    for title, (first, last) in SHEETS.items():
        scored = fill_sheet(wb[title], first, last, results["scores"][title], evaluators)
        if not scored:
            wrap_iferror(wb[title]["F1"])
            wrap_iferror(summary[f"D{SUMMARY_ROW[title]}"])
            print(f"note: every statement in '{title}' was N/A; sheet excluded from "
                  f"the global score")
    wrap_iferror(summary["E3"])
    wb.save(output)
    print(f"wrote {output}")


if __name__ == "__main__":
    main()
